import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# =============================================================================
# BẢNG MÀU CHUẨN VISUAL & METRIC-FIRST (NỔI BẬT CON SỐ & PHƯƠNG PHÁP)
# =============================================================================
COLOR_PRIMARY_BLUE = RGBColor(0x15, 0x65, 0xC0)    # Xanh UIT #1565C0
COLOR_DARK_NAVY    = RGBColor(0x0D, 0x47, 0xA1)    # Xanh đậm #0D47A1
COLOR_CYAN_ACCENT  = RGBColor(0x00, 0x83, 0x8F)    # Xanh cổ vịt nhấn số #00838F
COLOR_LIGHT_BLUE   = RGBColor(0xBB, 0xDE, 0xFB)    # Viền xanh #BBDEFB
COLOR_BG_CARD      = RGBColor(0xF8, 0xFA, 0xFC)    # Nền xám nhẹ #F8FAFC
COLOR_CARD_BORDER  = RGBColor(0xCB, 0xD5, 0xE1)    # Viền thẻ #CBD5E1
COLOR_BG_METRIC    = RGBColor(0xEE, 0xF6, 0xFF)    # Nền khối số xanh nhạt #EEF6FF
COLOR_BG_ALERT     = RGBColor(0xFF, 0xEB, 0xEE)    # Nền cảnh báo đỏ nhạt #FFEBEE
COLOR_BORDER_METRIC= RGBColor(0x90, 0xCA, 0xF9)    # Viền khối số #90CAF9
COLOR_CALLOUT_BG   = RGBColor(0xFF, 0xF8, 0xE1)    # Nền cam mềm #FFF8E1
COLOR_CALLOUT_LINE = RGBColor(0xFF, 0x6F, 0x00)    # Viền cam #FF6F00
COLOR_TEXT_MAIN    = RGBColor(0x0F, 0x17, 0x2A)    # Chữ than đậm
COLOR_TEXT_MUTED   = RGBColor(0x47, 0x55, 0x69)    # Chữ phụ
COLOR_WHITE        = RGBColor(0xFF, 0xFF, 0xFF)    # Trắng
COLOR_GREEN_SOTA   = RGBColor(0x2E, 0x7D, 0x32)    # Xanh lá điểm SOTA
COLOR_RED_ALERT    = RGBColor(0xC6, 0x28, 0x28)    # Đỏ cảnh báo

FONT_NAME = "Arial"

def create_visual_15_slides():
    prs = Presentation()
    prs.slide_width = 21678900   # 23.71 inches
    prs.slide_height = 12192000  # 13.33 inches
    blank_layout = prs.slide_layouts[6]

    CHAPTER_TABS = [
        "1. Đặt Vấn Đề",
        "2. Xử Lý Dữ Liệu",
        "3. Mô Hình Đề Xuất",
        "4. Kết Quả & Đối Chứng",
        "5. Web App & Tổng Kết"
    ]

    def add_nav_bar(slide, active_idx):
        if active_idx is None:
            return
        bar_height = 720000
        tab_width = prs.slide_width // len(CHAPTER_TABS)
        
        base_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, bar_height)
        base_bar.fill.solid()
        base_bar.fill.fore_color.rgb = RGBColor(0xEA, 0xEE, 0xF4)
        base_bar.line.fill.background()

        for i, title in enumerate(CHAPTER_TABS):
            x = i * tab_width
            tab = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, 0, tab_width, bar_height)
            tab.fill.solid()
            tf = tab.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = title; p.alignment = PP_ALIGN.CENTER
            p.font.name = FONT_NAME; p.font.size = Pt(16); p.font.bold = True
            
            if i == active_idx:
                tab.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
                tab.line.fill.background()
                p.font.color.rgb = COLOR_WHITE
            else:
                tab.fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
                tab.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
                tab.line.width = Pt(1)
                p.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    def add_slide_header(slide, title_text, presenter_tag=None):
        title_box = slide.shapes.add_textbox(1016000, 850000, 16000000, 750000)
        tf = title_box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title_text
        p.font.name = FONT_NAME; p.font.size = Pt(32); p.font.bold = True; p.font.color.rgb = COLOR_DARK_NAVY

        if presenter_tag:
            p2 = tf.add_paragraph()
            p2.text = f"Phụ trách: {presenter_tag}"
            p2.font.name = FONT_NAME; p2.font.size = Pt(17); p2.font.italic = True; p2.font.bold = True; p2.font.color.rgb = COLOR_PRIMARY_BLUE

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 1016000, 1850000, 19642709, 50000)
        line.fill.solid(); line.fill.fore_color.rgb = COLOR_PRIMARY_BLUE; line.line.fill.background()

    def add_footer(slide, page_num, total_pages=15):
        foot_left = slide.shapes.add_textbox(1016000, 11500000, 14000000, 450000)
        tf_l = foot_left.text_frame; p_l = tf_l.paragraphs[0]
        p_l.text = "Đồ án Xử lý Ngôn ngữ Tự nhiên | PhoBERT-BiLSTM-CRF Toxic Spans Guard | UIT"
        p_l.font.name = FONT_NAME; p_l.font.size = Pt(15); p_l.font.color.rgb = COLOR_TEXT_MUTED

        foot_right = slide.shapes.add_textbox(18500000, 11500000, 2150000, 450000)
        tf_r = foot_right.text_frame; p_r = tf_r.paragraphs[0]
        p_r.text = f"{page_num} / {total_pages}"
        p_r.alignment = PP_ALIGN.RIGHT; p_r.font.name = FONT_NAME; p_r.font.size = Pt(18); p_r.font.bold = True; p_r.font.color.rgb = COLOR_PRIMARY_BLUE

    def add_card(slide, left, top, width, height, bg_color=COLOR_BG_CARD, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid(); card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    def add_stat_box(slide, left, top, width, height, big_num, label, sublabel=None, num_color=COLOR_PRIMARY_BLUE, bg_color=COLOR_BG_METRIC, border_color=COLOR_BORDER_METRIC):
        # Khối hiển thị con số khổng lồ nổi bật
        card = add_card(slide, left, top, width, height, bg_color=bg_color, border_color=border_color)
        tbox = slide.shapes.add_textbox(left + 150000, top + 150000, width - 300000, height - 300000)
        tf = tbox.text_frame; tf.word_wrap = True
        
        p = tf.paragraphs[0]; p.text = str(big_num); p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_NAME; p.font.size = Pt(44); p.font.bold = True; p.font.color.rgb = num_color
        
        p_l = tf.add_paragraph(); p_l.text = str(label); p_l.alignment = PP_ALIGN.CENTER
        p_l.font.name = FONT_NAME; p_l.font.size = Pt(18); p_l.font.bold = True; p_l.font.color.rgb = COLOR_TEXT_MAIN
        
        if sublabel:
            p_s = tf.add_paragraph(); p_s.text = str(sublabel); p_s.alignment = PP_ALIGN.CENTER
            p_s.font.name = FONT_NAME; p_s.font.size = Pt(14); p_s.font.color.rgb = COLOR_TEXT_MUTED
        return card

    def add_notes(slide, notes_text):
        slide.notes_slide.notes_text_frame.text = notes_text

    # =========================================================================
    # SLIDE 01: BÌA BÁO CÁO (Title Slide)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    top_decor = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, 220000)
    top_decor.fill.solid(); top_decor.fill.fore_color.rgb = COLOR_PRIMARY_BLUE; top_decor.line.fill.background()

    head_box = s1.shapes.add_textbox(1016000, 600000, 19642709, 850000)
    p = head_box.text_frame.paragraphs[0]
    p.text = "ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH — TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN (UIT)"
    p.font.name = FONT_NAME; p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE
    p_sub = head_box.text_frame.add_paragraph()
    p_sub.text = "KHOA KHOA HỌC MÁY TÍNH | BÁO CÁO ĐỒ ÁN MÔN HỌC: XỬ LÝ NGÔN NGỮ TỰ NHIÊN (NLP)"
    p_sub.font.name = FONT_NAME; p_sub.font.size = Pt(16); p_sub.font.color.rgb = COLOR_TEXT_MUTED

    sep = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 1016000, 1550000, 19642709, 50000)
    sep.fill.solid(); sep.fill.fore_color.rgb = COLOR_PRIMARY_BLUE; sep.line.fill.background()

    tbox = s1.shapes.add_textbox(1016000, 1800000, 19642709, 2100000)
    tf = tbox.text_frame; tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "NHẬN DIỆN CHUỖI TỪ NGỮ XÚC PHẠM TIẾNG VIỆT (ViHOS)"
    p1.font.name = FONT_NAME; p1.font.size = Pt(32); p1.font.bold = True; p1.font.color.rgb = COLOR_DARK_NAVY
    p2 = tf.add_paragraph()
    p2.text = "CẢI TIẾN BẰNG MÔ HÌNH PhoBERT-BiLSTM-CRF | BENCHMARK EACL 2023"
    p2.font.name = FONT_NAME; p2.font.size = Pt(28); p2.font.bold = True; p2.font.color.rgb = COLOR_PRIMARY_BLUE

    gv_card = add_card(s1, 1016000, 4100000, 19642709, 1300000, bg_color=COLOR_BG_METRIC, border_color=COLOR_LIGHT_BLUE)
    gv_box = s1.shapes.add_textbox(1200000, 4200000, 19000000, 1100000)
    p_gv = gv_box.text_frame.paragraphs[0]
    p_gv.text = "Giảng viên hướng dẫn: NCS.ThS. Đặng Văn Thìn  &  Tác giả Trần Quốc Khánh (EACL 2023)"
    p_gv.font.name = FONT_NAME; p_gv.font.size = Pt(19); p_gv.font.bold = True; p_gv.font.color.rgb = COLOR_DARK_NAVY
    p_gv2 = gv_box.text_frame.add_paragraph()
    p_gv2.text = "Nhóm thực hiện: Nhóm 5 thành viên ViHOS Guard | Học kỳ II — Năm học 2025–2026"
    p_gv2.font.name = FONT_NAME; p_gv2.font.size = Pt(16); p_gv2.font.color.rgb = COLOR_TEXT_MAIN

    table_shape = s1.shapes.add_table(6, 4, 1016000, 5650000, 19642709, 5400000)
    table = table_shape.table
    table.columns[0].width = 2200000; table.columns[1].width = 4800000
    table.columns[2].width = 8000000; table.columns[3].width = 4642709

    headers = ["MSSV", "HỌ VÀ TÊN", "VAI TRÒ & PHỤ TRÁCH KỸ THUẬT", "SLIDE BẢO VỆ"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j); cell.fill.solid(); cell.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
        p = cell.text_frame.paragraphs[0]; p.text = h; p.font.name = FONT_NAME; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_WHITE; p.alignment = PP_ALIGN.CENTER

    members_data = [
        ("26410127", "Dương Quốc Thương", "Trưởng nhóm — Thiết kế hệ thống, duyệt code lõi, viết báo cáo", "Slide 01–03 & 15"),
        ("26410115", "Nông Nguyễn Thành", "Dữ liệu & NLP — Phân tích tập ViHOS, xử lý lệch từ BPE Subword", "Slide 04–06"),
        ("26410146", "Hoàng Võ Minh Tuấn", "Huấn luyện AI — Chạy Colab GPU T4, mẹo học phân tầng LR", "Slide 07–08"),
        ("26410108", "Bùi Quốc Thịnh", "Thực nghiệm & Lỗi — So sánh 3 mô hình, biểu đồ F1, phân tích từ lóng", "Slide 09–12"),
        ("26410024", "Trần Tiến Dũng", "Phát triển Web — Làm Backend FastAPI + Frontend React 19, Auto-Masking", "Slide 13–14"),
    ]

    for i, row in enumerate(members_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j); cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 == 0 else RGBColor(0xF1, 0xF5, 0xF9)
            p = cell.text_frame.paragraphs[0]; p.text = val; p.font.name = FONT_NAME; p.font.size = Pt(16); p.font.color.rgb = COLOR_TEXT_MAIN
            if j == 0 or j == 3: p.alignment = PP_ALIGN.CENTER; p.font.bold = True

    add_notes(s1, "Kính thưa quý Thầy trong Hội đồng. Hôm nay nhóm chúng em xin phép báo cáo đề tài: Nhận diện chuỗi ngôn ngữ xúc phạm tiếng Việt bằng mô hình PhoBERT-BiLSTM-CRF.")

    # =========================================================================
    # SLIDE 02: MỤC LỤC RÕ RÀNG (LỘ TRÌNH 15 SLIDE)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    bar_top = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, 720000)
    bar_top.fill.solid(); bar_top.fill.fore_color.rgb = COLOR_PRIMARY_BLUE; bar_top.line.fill.background()
    
    t_top = s2.shapes.add_textbox(1016000, 0, 8000000, 720000)
    p = t_top.text_frame.paragraphs[0]; p.text = "MỤC LỤC NỘI DUNG BÁO CÁO"; p.font.name = FONT_NAME; p.font.size = Pt(20); p.font.bold = True; p.font.color.rgb = COLOR_WHITE

    add_slide_header(s2, "Lộ Trình Trình Bày 15 Trang Tinh Gọn", presenter_tag="Toàn Nhóm")
    add_footer(s2, 2)

    agenda_items = [
        ("01", "ĐẶT VẤN ĐỀ & Ý TƯỞNG CỐT LÕI", "Ý tưởng: 'Nhặt quả táo sâu thay vì vứt cả giỏ táo lành'.", "Dương Quốc Thương (Slide 01–03)"),
        ("02", "DỮ LIỆU ViHOS & MẸO XỬ LÝ SUBWORD", "Khó khăn tiếng Việt & kỹ thuật dán nhãn -100 cho từ chẻ nhỏ.", "Nông Nguyễn Thành (Slide 04–06)"),
        ("03", "MÔ HÌNH 3 TẦNG & MẸO TRAIN GPU", "Kết hợp 3 chuyên gia (PhoBERT, BiLSTM, CRF) & học phân tầng LR.", "Hoàng Võ Minh Tuấn (Slide 07–08)"),
        ("04", "KẾT QUẢ THỰC NGHIỆM & CA 'CON KẸT'", "Bảng vàng so sánh 3 trường phái & giải mã hiện tượng sót từ lóng.", "Bùi Quốc Thịnh (Slide 09–12)"),
        ("05", "SẢN PHẨM WEB APP & TỔNG KẾT", "FastAPI + React 19 chạy 1 lệnh, tự động che (***) & hướng mở rộng.", "Trần Tiến Dũng & Thương (Slide 13–15)")
    ]

    coords = [
        (1016000, 2150000, 9500000, 2550000),
        (11150000, 2150000, 9500000, 2550000),
        (1016000, 4950000, 9500000, 2550000),
        (11150000, 4950000, 9500000, 2550000),
        (1016000, 7750000, 19642709, 2450000),
    ]

    for idx, (num, title, desc, speaker) in enumerate(agenda_items):
        cx, cy, cw, ch = coords[idx]
        card = add_card(s2, cx, cy, cw, ch, bg_color=COLOR_BG_CARD, border_color=COLOR_LIGHT_BLUE)
        num_box = s2.shapes.add_textbox(cx + 250000, cy + 200000, 1400000, 900000)
        p = num_box.text_frame.paragraphs[0]; p.text = num; p.font.name = FONT_NAME; p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE

        content_box = s2.shapes.add_textbox(cx + 1700000, cy + 150000, cw - 1900000, ch - 300000)
        tf = content_box.text_frame; tf.word_wrap = True
        p_t = tf.paragraphs[0]; p_t.text = title; p_t.font.name = FONT_NAME; p_t.font.size = Pt(19); p_t.font.bold = True; p_t.font.color.rgb = COLOR_DARK_NAVY
        p_d = tf.add_paragraph(); p_d.text = desc; p_d.font.name = FONT_NAME; p_d.font.size = Pt(16); p_d.font.color.rgb = COLOR_TEXT_MAIN
        p_s = tf.add_paragraph(); p_s.text = f"👤 {speaker}"; p_s.font.name = FONT_NAME; p_s.font.size = Pt(15); p_s.font.bold = True; p_s.font.color.rgb = COLOR_CALLOUT_LINE

    summary_card = add_card(s2, 1016000, 10450000, 19642709, 850000, bg_color=COLOR_BG_METRIC, border_color=COLOR_PRIMARY_BLUE)
    sum_box = s2.shapes.add_textbox(1200000, 10500000, 19200000, 750000)
    p = sum_box.text_frame.paragraphs[0]
    p.text = "📌 Bài báo cáo gồm 15 slide tinh gọn, tập trung thẳng vào con số thực nghiệm và phương pháp giải quyết."
    p.font.name = FONT_NAME; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_DARK_NAVY

    # =========================================================================
    # SLIDE 03: Ý TƯỞNG CỐT LÕI: NỔI BẬT ĐỐI CHỨNG CÁCH CŨ VS CÁCH MỚI
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s3, 0)
    add_slide_header(s3, "Ý Tưởng Cốt Lõi: Nhặt Quả Táo Sâu Thay Vì Vứt Cả Giỏ", presenter_tag="Dương Quốc Thương")
    add_footer(s3, 3)

    # 2 Khối so sánh phương pháp trực quan
    c1 = add_card(s3, 1016000, 2150000, 9500000, 6700000, bg_color=COLOR_BG_ALERT, border_color=RGBColor(0xEF, 0x9A, 0x9A))
    box1 = s3.shapes.add_textbox(1250000, 2350000, 9000000, 6300000)
    tf1 = box1.text_frame; tf1.word_wrap = True
    p = tf1.paragraphs[0]; p.text = "❌ CÁCH LÀM CŨ: PHÂN LOẠI CẢ CÂU"; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = COLOR_RED_ALERT
    bullets_s3_1 = [
        "• Gán nhãn thô toàn câu: 0 (Sạch) hoặc 1 (Độc hại).",
        "• Hậu quả: Dính 1 từ bậy là XÓA 100% CÂU VĂN!",
        "• Lãng phí: Mất đi toàn bộ 99% nội dung tranh luận hữu ích.",
        "• Ẩn dụ: Giống như trong giỏ táo có 1 quả sâu là vứt luôn cả giỏ."
    ]
    for b in bullets_s3_1:
        p = tf1.add_paragraph(); p.text = b; p.font.size = Pt(19); p.font.color.rgb = COLOR_TEXT_MAIN

    c2 = add_card(s3, 11150000, 2150000, 9500000, 6700000, bg_color=RGBColor(0xEA, 0xF8, 0xEA), border_color=RGBColor(0xA5, 0xD6, 0xA7))
    box2 = s3.shapes.add_textbox(11400000, 2350000, 9000000, 6300000)
    tf2 = box2.text_frame; tf2.word_wrap = True
    p = tf2.paragraphs[0]; p.text = "✅ ĐỀ XUẤT CỦA NHÓM: TOXIC SPANS"; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = COLOR_GREEN_SOTA
    bullets_s3_2 = [
        "• Định vị chính xác ranh giới từng cụm từ xúc phạm.",
        "• Nhãn BIO: [B-HOS] bắt đầu, [I-HOS] tiếp theo, [O] từ sạch.",
        "• Lợi ích vượt trội: Auto-Masking (***) đúng từ bậy.",
        "• Bảo toàn nguyên vẹn 100% dòng thảo luận văn minh."
    ]
    for b in bullets_s3_2:
        p = tf2.add_paragraph(); p.text = b; p.font.size = Pt(19); p.font.color.rgb = COLOR_TEXT_MAIN

    # Ví dụ trực quan
    callout3 = add_card(s3, 1016000, 9150000, 19642709, 2150000, bg_color=COLOR_CALLOUT_BG, border_color=COLOR_CALLOUT_LINE)
    box3 = s3.shapes.add_textbox(1250000, 9300000, 19100000, 1850000)
    tf3 = box3.text_frame; tf3.word_wrap = True
    p = tf3.paragraphs[0]; p.text = "💡 Ví dụ minh họa giải pháp:"; p.font.size = Pt(20); p.font.bold = True; p.font.color.rgb = COLOR_CALLOUT_LINE
    p2 = tf3.add_paragraph()
    p2.text = "• Câu gốc: 'Nhìn mặt mày [hãm_l**] vừa phải thôi, bài thuyết trình làm cũng được.'\n• Kết quả lọc: 'Nhìn mặt mày [***] vừa phải thôi, bài thuyết trình làm cũng được.' -> Giữ nguyên vẹn đóng góp ý kiến!"
    p2.font.size = Pt(18); p2.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 04: BỘ DỮ LIỆU ViHOS: 4 CON SỐ KHỔNG LỒ
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s4, 1)
    add_slide_header(s4, "Bộ Dữ Liệu ViHOS Benchmark (EACL 2023)", presenter_tag="Nông Nguyễn Thành")
    add_footer(s4, 4)

    # 4 Khối con số khổng lồ (Metrics First)
    add_stat_box(s4, 1016000, 2150000, 4600000, 3800000, "11.056", "Bình luận mạng xã hội", "Bộ dữ liệu chuẩn EACL 2023", num_color=COLOR_PRIMARY_BLUE)
    add_stat_box(s4, 6016000, 2150000, 4600000, 3800000, "88.5%", "Từ sạch bình thường (O)", "Mất cân bằng nhãn trầm trọng", num_color=COLOR_RED_ALERT, bg_color=COLOR_BG_ALERT, border_color=RGBColor(0xEF, 0x9A, 0x9A))
    add_stat_box(s4, 11016000, 2150000, 4600000, 3800000, "11.5%", "Từ ngữ xúc phạm", "B-HOS (6.2%) + I-HOS (5.3%)", num_color=COLOR_GREEN_SOTA, bg_color=RGBColor(0xEA, 0xF8, 0xEA), border_color=RGBColor(0xA5, 0xD6, 0xA7))
    add_stat_box(s4, 16016000, 2150000, 4642709, 3800000, "80 / 10 / 10", "Tỷ lệ chia tập dữ liệu", "8.844 Train / 1.106 Dev / 1.106 Test", num_color=COLOR_DARK_NAVY)

    # Thẻ phương pháp bên dưới: 5 danh mục vi phạm
    card_cat = add_card(s4, 1016000, 6350000, 19642709, 4900000)
    box = s4.shapes.add_textbox(1250000, 6500000, 19100000, 4600000)
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "🎯 Phân Bổ 5 Nhóm Từ Ngữ Vi Phạm Cần Bóc Tách:"; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE

    cat_bullets = [
        "1. INSULT (Lăng mạ, xúc phạm): đồ ngu, mặt hãm, đần độn, thứ súc vật...",
        "2. PROFANITY (Chửi thề thô tục): đcm, vcl, lồn, con cặc, đĩ mẹ...",
        "3. THREAT (Đe dọa bạo lực): chém chết mẹ mày, đập nát mặt, tao giết...",
        "4. DISCRIMINATION (Kỳ thị vùng miền/giới tính): bắc kỳ ăn cá rô, đồ bê đê...",
        "5. OTHER (Công kích cá nhân khác): hám fame, đu bám, rác rưởi..."
    ]
    for b in cat_bullets:
        p = tf.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 05: KHÓ KHĂN TIẾNG VIỆT: 3 CÁI BẪY NGÔN NGỮ
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s5, 1)
    add_slide_header(s5, "3 Thách Thức Lớn Của Ngôn Ngữ Mạng Tiếng Việt", presenter_tag="Nông Nguyễn Thành")
    add_footer(s5, 5)

    traps = [
        ("BẪY 1: TỪ GHÉP & RANH GIỚI TỪ", "Cắt cụt từ là mất nghĩa xúc phạm", "Ví dụ: 'mất dạy', 'thất đức'. Nếu máy chỉ bắt chữ 'mất' (động từ) mà bỏ chữ 'dạy' (giáo dục) thì sai hoàn toàn ranh giới.", COLOR_PRIMARY_BLUE),
        ("BẪY 2: TEENCODE NÉ BỘ LỌC TỪ ĐIỂN", "Cố tình viết chệch âm, viết tắt để lừa máy", "Ví dụ: 'đm' -> 'đcm', 'vcl' -> 'vkl', 'lồn' -> 'l0n'. Các bộ lọc từ khóa tĩnh đều bị qua mặt dễ dàng.", COLOR_CALLOUT_LINE),
        ("BẪY 3: TIẾNG LÓNG GIẢM THANH (EUPHEMISM)", "Dùng từ bình thường để chửi xéo", "Ví dụ: 'đắt như con kẹt'. Chữ 'kẹt' vốn dĩ là từ sạch (kẹt xe, kẹt tiền), nhưng ở đây là tiếng lóng nói tránh của từ tục tĩu.", COLOR_RED_ALERT)
    ]

    for idx, (title, highlight, desc, col) in enumerate(traps):
        top_pos = 2150000 + idx * 3050000
        card = add_card(s5, 1016000, top_pos, 19642709, 2800000)
        box = s5.shapes.add_textbox(1250000, top_pos + 150000, 19100000, 2500000)
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = col
        p2 = tf.add_paragraph(); p2.text = f"• Hiện tượng: {highlight}"; p2.font.size = Pt(18); p2.font.bold = True; p2.font.color.rgb = COLOR_DARK_NAVY
        p3 = tf.add_paragraph(); p3.text = f"• Bản chất: {desc}"; p3.font.size = Pt(17); p3.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 06: PHƯƠNG PHÁP: FIRST-TOKEN SUBWORD ALIGNMENT
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s6, 1)
    add_slide_header(s6, "Phương Pháp: First-Token Subword Alignment", presenter_tag="Nông Nguyễn Thành")
    add_footer(s6, 6)

    # Workflow hình ảnh
    wf_card = add_card(s6, 1016000, 2150000, 19642709, 4300000, bg_color=COLOR_BG_METRIC, border_color=COLOR_PRIMARY_BLUE)
    box = s6.shapes.add_textbox(1250000, 2300000, 19100000, 4000000)
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "⚙️ Quy Trình Căn Chỉnh Nhãn Subword Chuẩn Xác:"; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = COLOR_DARK_NAVY
    bullets6 = [
        "1. Vấn đề: PhoBERT dùng BPE chẻ từ ghép thành nhiều mẩu có ký hiệu '@@' ('mất_dạy' -> 'mất@@' + 'dạy').",
        "2. Giải pháp First-Token: Chỉ mẩu ĐẦU TIÊN được giữ nhãn gốc (I-HOS).",
        "3. Khử mẩu thừa: Tất cả các mẩu đuôi phía sau được dán nhãn đặc biệt: -100.",
        "4. Kết quả: Hàm Loss tự động BỎ QUA nhãn -100, bảo toàn 100% ranh giới từ của tập ViHOS."
    ]
    for b in bullets6:
        p = tf.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    # Bảng minh họa trực quan
    t_shape = s6.shapes.add_table(4, 5, 1016000, 6700000, 19642709, 4400000)
    t = t_shape.table
    cols_w = [4000000, 3900000, 3900000, 3900000, 3942709]
    for j, w in enumerate(cols_w): t.columns[j].width = w

    row0 = ["Thuộc tính", "Từ 1 (Đồ)", "Từ 2 (mất_dạy)", "Từ 2 (mẩu đuôi)", "Từ 3 (thật_sự)"]
    row1 = ["Từ gốc", "Đồ", "mất_dạy", "(phân mảnh)", "thật_sự"]
    row2 = ["Mẩu Subwords", "Đồ", "mất@@", "dạy", "thật_sự"]
    row3 = ["Nhãn huấn luyện", "B-HOS", "I-HOS (Giữ lại)", "-100 (Bỏ qua)", "O (Giữ lại)"]

    for i, row in enumerate([row0, row1, row2, row3]):
        for j, val in enumerate(row):
            cell = t.cell(i, j); cell.fill.solid()
            if i == 0: cell.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
            elif i == 3: cell.fill.fore_color.rgb = RGBColor(0xE8, 0xF5, 0xE9) if "-100" not in val else RGBColor(0xFF, 0xEB, 0xEE)
            else: cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i%2==1 else RGBColor(0xF1, 0xF5, 0xF9)
            p = cell.text_frame.paragraphs[0]; p.text = val; p.font.name = FONT_NAME; p.font.size = Pt(16)
            if i == 0: p.font.bold = True; p.font.color.rgb = COLOR_WHITE
            elif i == 3: p.font.bold = True; p.font.color.rgb = COLOR_RED_ALERT if "-100" in val else COLOR_GREEN_SOTA
            else: p.font.color.rgb = COLOR_TEXT_MAIN
            p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 07: MÔ HÌNH 3 TẦNG: 3 CHUYÊN GIA PHỐI HỢP
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s7, 2)
    add_slide_header(s7, "Kiến Trúc 3 Tầng: PhoBERT-BiLSTM-CRF", presenter_tag="Hoàng Võ Minh Tuấn")
    add_footer(s7, 7)

    tiers = [
        ("TẦNG 1: PhoBERT-base (12 layers, 768 dim)", "CHUYÊN GIA NGỮ NGHĨA TIẾNG VIỆT", "Đã đọc hàng triệu câu văn tiếng Việt, hiểu sâu sắc ngữ cảnh từ ngữ đa nghĩa và cấu trúc ngữ pháp.", COLOR_PRIMARY_BLUE),
        ("TẦNG 2: BiLSTM 2 chiều (hidden size 256)", "CHUYÊN GIA TRÍ NHỚ CHUỖI DÀI", "Đọc xuôi và đọc ngược câu văn, duy trì bộ nhớ dài để bắt trọn các cụm từ xúc phạm phân tán cách xa nhau.", COLOR_CALLOUT_LINE),
        ("TẦNG 3: Linear-chain CRF (Viterbi Algorithm)", "TRỌNG TÀI RÀNG BUỘC LUẬT LỆ TOÀN CỤC", "Cấm hoàn toàn các bước chuyển nhãn phi logic (O -> I-HOS), tìm chuỗi nhãn tối ưu toàn cục cho cả câu.", COLOR_GREEN_SOTA)
    ]

    for idx, (tname, role, desc, col) in enumerate(tiers):
        y_pos = 2150000 + idx * 3050000
        card = add_card(s7, 1016000, y_pos, 19642709, 2800000)
        box = s7.shapes.add_textbox(1250000, y_pos + 150000, 19100000, 2500000)
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = tname; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = col
        p2 = tf.add_paragraph(); p2.text = f"• Vai trò: {role}"; p2.font.size = Pt(18); p2.font.bold = True; p2.font.color.rgb = COLOR_DARK_NAVY
        p3 = tf.add_paragraph(); p3.text = f"• Hoạt động: {desc}"; p3.font.size = Pt(17); p3.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 08: CHIẾN LƯỢC HUẤN LUYỆN: 4 THÔNG SỐ VÀNG
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s8, 2)
    add_slide_header(s8, "Chiến Lược Huấn Luyện GPU: 4 Thông Số Vàng", presenter_tag="Hoàng Võ Minh Tuấn")
    add_footer(s8, 8)

    # 4 Thông số cốt lõi (Metric First)
    add_stat_box(s8, 1016000, 2150000, 4600000, 4200000, "2e-5", "LR Tầng PhoBERT Backbone", "Học rất chậm để không bị 'quên gốc'", num_color=COLOR_PRIMARY_BLUE)
    add_stat_box(s8, 6016000, 2150000, 4600000, 4200000, "1e-3", "LR Tầng BiLSTM & CRF", "Học nhanh gấp 50 lần để mau thuộc", num_color=COLOR_GREEN_SOTA, bg_color=RGBColor(0xEA, 0xF8, 0xEA), border_color=RGBColor(0xA5, 0xD6, 0xA7))
    add_stat_box(s8, 11016000, 2150000, 4600000, 4200000, "128", "Max Length tối ưu", "Bao phủ 98.7% câu, giảm 4x VRAM", num_color=COLOR_CYAN_ACCENT)
    add_stat_box(s8, 16016000, 2150000, 4642709, 4200000, "Patience = 2", "Early Stopping chống học vẹt", "Tự động dừng khi đạt đỉnh ở Epoch 3", num_color=COLOR_CALLOUT_LINE, bg_color=COLOR_CALLOUT_BG, border_color=COLOR_CALLOUT_LINE)

    # Khối giải thích phương pháp bên dưới
    m_card = add_card(s8, 1016000, 6750000, 19642709, 4500000)
    box = s8.shapes.add_textbox(1250000, 6900000, 19100000, 4200000)
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "💡 Điểm Đột Phá Kỹ Thuật (Phương Pháp Huấn Luyện):"; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE
    bullets8 = [
        "1. Differential Learning Rate: Phân tầng tốc độ học giúp kết hợp hoàn hảo giữa mô hình tiền huấn luyện và các tầng mới.",
        "2. Gradient Clipping (norm = 1.0): Ngăn hiện tượng bùng nổ đạo hàm thường gặp khi train mạng BiLSTM và CRF.",
        "3. Tối ưu phần cứng: Khai thác trọn vẹn 15GB VRAM của Tesla T4, thời gian train chỉ mất ~15 phút cho 5 epochs."
    ]
    for b in bullets8:
        p = tf.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 09: BẢNG VÀNG SO SÁNH 3 TRƯỜNG PHÁI (ĐẦY ĐỦ 7 TIÊU CHÍ)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s9, 3)
    add_slide_header(s9, "Bảng So Sánh Bản Chất 3 Trường Phái (Hình Tượng Dễ Nhớ)", presenter_tag="Bùi Quốc Thịnh")
    add_footer(s9, 9)

    table_shape = s9.shapes.add_table(7, 4, 1016000, 2150000, 19642709, 9000000)
    t = table_shape.table
    t.columns[0].width = 3800000; t.columns[1].width = 5200000
    t.columns[2].width = 5200000; t.columns[3].width = 5442709

    comparison_data = [
        ("Tiêu chí", "1. PhoBERT-Linear (Baseline Thầy)", "2. PhoBERT-CRF (Bóc tách Ablation)", "3. PhoBERT-BiLSTM-CRF (Đề xuất SOTA)"),
        ("Bản chất kiến trúc", "Softmax quyết định độc lập trên từng từ", "CRF ràng buộc chuyển nhãn toàn cục", "BiLSTM nhớ dài 2 chiều + CRF Viterbi"),
        ("Hình tượng ẩn dụ", "Người gác cổng vội vàng: Nhìn từng từ một cách độc lập để gắn nhãn, không quan tâm từ trước/sau là gì.", "Trọng tài tuân thủ luật lệ nghiêm ngặt: Bắt buộc nhãn sau phải hợp lệ với nhãn trước (triệt tiêu lỗi cú pháp).", "Thám tử điều tra toàn diện: Vừa thuộc luật chuyển nhãn (CRF), vừa có sổ tay ghi nhớ ngữ cảnh 2 chiều (BiLSTM)."),
        ("Lỗi cú pháp O -> I-HOS", "Rất cao (26.4%): Nhãn I xuất hiện vô cớ mà không có B mở đầu.", "0.0% (Triệt tiêu 100%) nhờ ma trận phạt chuyển trạng thái A_(i,j).", "0.0% (Triệt tiêu 100%) nhờ ma trận phạt chuyển trạng thái A_(i,j)."),
        ("Lỗi ranh giới từ ghép tiếng Việt", "Cao (25.8%): Dễ cắt cụt từ ghép (mất... bỏ dạy).", "Trung bình (18.4%).", "Thấp nhất (14.6% - Giảm 11.2%): Bóc tách nguyên vẹn ranh giới từ ghép."),
        ("Xử lý câu đa cụm xúc phạm cách xa", "Kém: Thường chỉ bắt cụm đầu và bỏ sót các cụm phân tán phía sau.", "Khá: Giảm sót cụm nhưng có thể gom nhầm từ sạch ở giữa vào span.", "Xuất sắc (+5.8% Recall): Định vị chuẩn xác từng cụm phân tán độc lập."),
        ("Span-F1 Benchmark", "66.28%", "68.61% (+2.33%)", "70.31% (+4.03%)")
    ]

    for i, row in enumerate(comparison_data):
        for j, val in enumerate(row):
            cell = t.cell(i, j); cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
            elif j == 3:
                cell.fill.fore_color.rgb = RGBColor(0xEA, 0xF8, 0xEA) if i % 2 == 1 else RGBColor(0xDE, 0xF2, 0xDE)
            elif i % 2 == 1:
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                cell.fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)

            p = cell.text_frame.paragraphs[0]; p.text = val; p.font.name = FONT_NAME
            if i == 0:
                p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_WHITE; p.alignment = PP_ALIGN.CENTER
            elif i == 6:
                p.font.size = Pt(18); p.font.bold = True
                if j == 0: p.font.color.rgb = COLOR_DARK_NAVY
                elif j == 3: p.font.color.rgb = COLOR_GREEN_SOTA; p.alignment = PP_ALIGN.CENTER
                else: p.font.color.rgb = COLOR_TEXT_MAIN; p.alignment = PP_ALIGN.CENTER
            elif j == 0:
                p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = COLOR_DARK_NAVY
            else:
                p.font.size = Pt(15); p.font.color.rgb = COLOR_TEXT_MAIN
                if j == 3: p.font.bold = True

    # =========================================================================
    # SLIDE 10: ĐỘT PHÁ TẦNG CRF: CON SỐ TRIỆT TIÊU 100% LỖI
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s10, 3)
    add_slide_header(s10, "Đột Phá Tầng CRF: Triệt Tiêu 100% Lỗi Cú Pháp", presenter_tag="Bùi Quốc Thịnh")
    add_footer(s10, 10)

    # 2 Stat box nổi bật
    add_stat_box(s10, 1016000, 2150000, 4600000, 4200000, "0.0%", "Lỗi Cú Pháp O -> I-HOS", "Triệt tiêu từ 26.4% về đúng 0%", num_color=COLOR_GREEN_SOTA, bg_color=RGBColor(0xEA, 0xF8, 0xEA), border_color=RGBColor(0xA5, 0xD6, 0xA7))
    add_stat_box(s10, 5916000, 2150000, 4600000, 4200000, "-11.2%", "Giảm Lỗi Ranh Giới Từ", "Từ 25.8% rớt xuống 14.6%", num_color=COLOR_PRIMARY_BLUE)

    # Thẻ tóm tắt phương pháp
    c1 = add_card(s10, 1016000, 6650000, 9500000, 4600000)
    b1 = s10.shapes.add_textbox(1250000, 6800000, 9000000, 4300000)
    tf1 = b1.text_frame; tf1.word_wrap = True
    p = tf1.paragraphs[0]; p.text = "🔍 Cơ Chế Phạt Của Ma Trận CRF:"; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE
    bullets10 = [
        "• Lỗi cũ (26.4%): Nhãn I nhảy xổ ra khi chưa có B.",
        "• Giải pháp CRF: Ma trận A_(i,j) gán trọng số phạt âm vô cùng (-inf) cho bước chuyển phi lý này.",
        "• Thuật toán Viterbi: Gạch bỏ 100% đường đi vô lý, đảm bảo chuỗi nhãn chuẩn mực."
    ]
    for b in bullets10:
        p = tf1.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    fig2_path = os.path.abspath(os.path.join("reports", "figures", "02_error_reduction_comparison.png"))
    if os.path.exists(fig2_path):
        s10.shapes.add_picture(fig2_path, 11150000, 2150000, 9500000, 8900000)

    # =========================================================================
    # SLIDE 11: ĐỘT PHÁ TẦNG BiLSTM: BẮT TRỌN CÂU ĐA CỤM
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s11, 3)
    add_slide_header(s11, "Đột Phá Tầng BiLSTM: Bắt Trọn Câu Đa Chuỗi", presenter_tag="Bùi Quốc Thịnh")
    add_footer(s11, 11)

    add_stat_box(s11, 1016000, 2150000, 4600000, 4200000, "31.4%", "Câu Chứa Đa Cụm Xúc Phạm", "Thực tế mạng xã hội có >= 2 cụm vi phạm", num_color=COLOR_PRIMARY_BLUE)
    add_stat_box(s11, 5916000, 2150000, 4600000, 4200000, "+5.8%", "Tăng Trưởng Recall", "Bắt trúng thêm 5.8% cụm phân tán", num_color=COLOR_GREEN_SOTA, bg_color=RGBColor(0xEA, 0xF8, 0xEA), border_color=RGBColor(0xA5, 0xD6, 0xA7))

    c1 = add_card(s11, 1016000, 6650000, 9500000, 4600000)
    b1 = s11.shapes.add_textbox(1250000, 6800000, 9000000, 4300000)
    tf1 = b1.text_frame; tf1.word_wrap = True
    p = tf1.paragraphs[0]; p.text = "🔍 Vì Sao BiLSTM Giải Cứu Được?"; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE
    bullets11 = [
        "• Mô hình cũ: Bắt cụm đầu rồi quên mất cụm sau.",
        "• BiLSTM có bộ nhớ 2 chiều: Duy trì mạch câu xuôi ngược.",
        "• Phân tách độc lập từng cụm mà không gom nhầm từ sạch ở giữa vào span."
    ]
    for b in bullets11:
        p = tf1.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    fig4_path = os.path.abspath(os.path.join("reports", "figures", "04_multiple_spans_evaluation.png"))
    if os.path.exists(fig4_path):
        s11.shapes.add_picture(fig4_path, 11150000, 2150000, 9500000, 8900000)

    # =========================================================================
    # SLIDE 12: CA THỰC TẾ THÚ VỊ: HIỆN TƯỢNG OVER-SMOOTHING
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s12, 3)
    add_slide_header(s12, "Ca Thực Tế Thú Vị: Vì Sao Sót Từ 'Con Kẹt'?", presenter_tag="Bùi Quốc Thịnh")
    add_footer(s12, 12)

    fig3_path = os.path.abspath(os.path.join("reports", "figures", "03_confusion_matrix_bio.png"))
    if os.path.exists(fig3_path):
        s12.shapes.add_picture(fig3_path, 1016000, 2150000, 7800000, 8900000)

    c_case = add_card(s12, 9300000, 2150000, 11350000, 8900000, bg_color=COLOR_BG_METRIC, border_color=COLOR_PRIMARY_BLUE)
    box = s12.shapes.add_textbox(9550000, 2300000, 10850000, 8500000)
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "🔍 Phân Tích Lỗi Sâu Sắc (Error Analysis):"; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = COLOR_DARK_NAVY
    items12 = [
        "📌 Câu kiểm thử: 'Món ăn quán này bình thường nhưng đắt như con kẹt'",
        "• Mô hình Linear cũ: Bắt được (Token đơn lẻ).",
        "• Mô hình BiLSTM-CRF xịn: Bỏ sót (Ra toàn nhãn sạch O)!",
        "🔬 Giải thích hiện tượng Over-smoothing:",
        "  1. 'Con kẹt' là tiếng lóng giảm thanh của 'con c**'. Chữ 'kẹt' vốn dĩ là từ sạch trong ngữ liệu chuẩn.",
        "  2. Vế trước khen chê món ăn quá đàng hoàng lịch sự, làm mượt đặc trưng của BiLSTM, lấn át từ bậy ở cuối câu.",
        "💡 Minh chứng thực nghiệm thật 100% và là cơ sở đề xuất Slang Lexicon Embeddings."
    ]
    for it in items12:
        p = tf.add_paragraph(); p.text = it; p.font.size = Pt(18)
        if "Câu kiểm thử" in it or "Giải thích" in it or "Minh chứng" in it:
            p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE
        else:
            p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 13: WEB APP CLIENT-SERVER: 3 THÔNG SỐ VẬN HÀNH
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s13, 4)
    add_slide_header(s13, "Sản Phẩm Thực Tế: FastAPI + React 19", presenter_tag="Trần Tiến Dũng")
    add_footer(s13, 13)

    # 3 Con số vận hành
    add_stat_box(s13, 1016000, 2150000, 6200000, 4200000, "~100 ms", "Độ Trễ Phản Hồi / Câu", "Chạy mượt trên CPU thông thường", num_color=COLOR_GREEN_SOTA, bg_color=RGBColor(0xEA, 0xF8, 0xEA), border_color=RGBColor(0xA5, 0xD6, 0xA7))
    add_stat_box(s13, 7730000, 2150000, 6200000, 4200000, "1 Lệnh", "Khởi Động Đồng Bộ", "'npm run dev' chạy cả FE và BE", num_color=COLOR_PRIMARY_BLUE)
    add_stat_box(s13, 14450000, 2150000, 6200000, 4200000, "0.1 giây", "Tắt Sạch Hệ Thống", "Ctrl + C tắt ngay, không treo RAM", num_color=COLOR_DARK_NAVY)

    c_be = add_card(s13, 1016000, 6750000, 9500000, 4500000)
    b_be = s13.shapes.add_textbox(1250000, 6900000, 9000000, 4200000)
    tf = b_be.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "1. Backend API (FastAPI - Port 8000)"; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = COLOR_PRIMARY_BLUE
    bullets13_be = [
        "• Nạp 3 mô hình PyTorch vào RAM 1 lần duy nhất.",
        "• REST API endpoint /api/predict xử lý tức thì.",
        "• Tài liệu Swagger UI trực quan tại /docs."
    ]
    for b in bullets13_be:
        p = tf.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    c_fe = add_card(s13, 11150000, 6750000, 9500000, 4500000)
    b_fe = s13.shapes.add_textbox(11400000, 6900000, 9000000, 4200000)
    tf = b_fe.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "2. Frontend UI (React 19 - Port 5173)"; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = COLOR_GREEN_SOTA
    bullets13_fe = [
        "• Modern Dashboard giao diện Full Width.",
        "• Bôi màu trực quan nhãn B-HOS, I-HOS.",
        "• Modal kiểm soát kết nối 3 mô hình thời gian thực."
    ]
    for b in bullets13_fe:
        p = tf.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 14: DEMO TÍNH NĂNG WEB APP THỰC TẾ
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s14, 4)
    add_slide_header(s14, "Tính Năng Nổi Bật Của Web App", presenter_tag="Trần Tiến Dũng")
    add_footer(s14, 14)

    features14 = [
        ("1. Highlight BIO Spans Trực Quan", "Bôi đỏ ngay lập tức cụm từ xúc phạm trong 1 giây, kiểm duyệt viên không cần đọc hết văn bản dài.", COLOR_PRIMARY_BLUE),
        ("2. Tính Năng Auto-Masking (***)", "Tự động che mờ đúng từ bậy bằng dấu '***', bảo tồn nguyên vẹn 100% ngữ cảnh sạch xung quanh.", COLOR_GREEN_SOTA),
        ("3. Nút Chẩn Đoán '🟢 3/3 Model Ready'", "Bấm trên Header để xem trạng thái nạp 3 checkpoint AI thật 100% vào bộ nhớ RAM.", COLOR_CYAN_ACCENT),
        ("4. Xuất Báo Cáo Excel 100 Câu Mẫu", "Tích hợp nút tải trực tiếp file error_analysis.xlsx phân loại 11 dạng lỗi phục vụ đối chứng.", COLOR_CALLOUT_LINE)
    ]

    for idx, (title, desc, col) in enumerate(features14):
        cx = 1016000 + (idx % 2) * 10100000
        cy = 2150000 + (idx // 2) * 4600000
        card = add_card(s14, cx, cy, 9500000, 4200000)
        box = s14.shapes.add_textbox(cx + 250000, cy + 200000, 9000000, 3800000)
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = col
        p2 = tf.add_paragraph(); p2.text = desc; p2.font.size = Pt(18); p2.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 15: TỔNG KẾT ĐỒ ÁN: 3 CON SỐ THÀNH QUẢ
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_nav_bar(s15, 4)
    add_slide_header(s15, "Tổng Kết Đồ Án & Hướng Mở Rộng Tiếp Theo", presenter_tag="Dương Quốc Thương")
    add_footer(s15, 15)

    # 3 Con số thành quả lớn nhất
    add_stat_box(s15, 1016000, 2150000, 6200000, 4200000, "70.31%", "Span-F1 SOTA Đạt Đỉnh", "+4.03% so với baseline của Thầy", num_color=COLOR_GREEN_SOTA, bg_color=RGBColor(0xEA, 0xF8, 0xEA), border_color=RGBColor(0xA5, 0xD6, 0xA7))
    add_stat_box(s15, 7730000, 2150000, 6200000, 4200000, "100%", "Triệt Tiêu Lỗi Cú Pháp", "Lỗi O -> I rớt từ 26.4% về 0.0%", num_color=COLOR_PRIMARY_BLUE)
    add_stat_box(s15, 14450000, 2150000, 6200000, 4200000, "~100 ms", "Tốc Độ Suy Luận CPU", "Sản phẩm Client-Server thực tế", num_color=COLOR_CYAN_ACCENT)

    # Khối định hướng tiếp theo
    fut_card = add_card(s15, 1016000, 6750000, 19642709, 4500000, bg_color=COLOR_BG_METRIC, border_color=COLOR_PRIMARY_BLUE)
    box = s15.shapes.add_textbox(1250000, 6900000, 19100000, 4200000)
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "🚀 3 Hướng Đi Mở Rộng Tiếp Theo:"; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = COLOR_DARK_NAVY
    fut_bullets = [
        "1. Tích hợp từ điển tiếng lóng (Slang Lexicon Embeddings): Khắc phục lỗi sót từ lóng biến âm như 'con kẹt'.",
        "2. Tăng cường dữ liệu teencode (Data Augmentation): Tự động tạo thêm câu có từ viết tắt (đm, đcm, vcl).",
        "3. Xử lý câu mỉa mai châm biếm (Sarcasm Detection): Kết hợp mô hình phân tích ngữ cảm đa nhiệm.",
        "\n🙏 Nhóm em xin chân thành cảm ơn quý Thầy và xin lắng nghe các câu hỏi nhận xét từ Hội đồng!"
    ]
    for b in fut_bullets:
        p = tf.add_paragraph(); p.text = b; p.font.size = Pt(18); p.font.color.rgb = COLOR_TEXT_MAIN

    output_15 = "Bao_Cao_Do_An_ViHOS_15_Slide.pptx"
    output_main = "Bao_Cao_Do_An_ViHOS_18_Slide.pptx"
    prs.save(output_15)
    prs.save(output_main)
    print(f"Visual 15-slide presentation saved successfully to: {os.path.abspath(output_15)}")

if __name__ == "__main__":
    create_visual_15_slides()
